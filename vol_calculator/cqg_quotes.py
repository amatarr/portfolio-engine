from openpyxl import load_workbook

# Header text CQGXL prints above each commodity's block of symbol/price
# rows in columns O/P, mapped to the commodity names this app uses.
_COMMODITY_HEADERS = {
    "CORN": "Corn",
    "SOYBEANS": "Soybean",
    "WHEAT": "SRW Wheat",
}


def _convert_cqg_symbol(symbol):
    """
    CQG's own symbols (e.g. "ZCEZ26", "ZSEX26", "ZWAZ26") carry one extra
    filler letter after the 2-letter root that this app's contract codes
    (e.g. "ZCZ26") don't use. All observed symbols are a fixed 6 characters
    -- root(2) + filler(1) + month(1) + year(2) -- so drop index 2.
    """

    s = symbol.strip()
    return s[:2] + s[3:]


def import_cqg_quotes(path):
    """
    Reads the CQGXL "RT links" export, columns O/P: a commodity header
    ("CORN"/"SOYBEANS"/"WHEAT") followed by rows of (symbol, price).

    Returns a list of {"commodity", "contract", "price"} dicts, contract
    already converted to this app's symbol convention.
    """

    wb = load_workbook(path, data_only=True)

    quotes = []
    current_commodity = None

    for ws in wb.worksheets:
        for o_cell, p_cell in ws.iter_rows(min_col=15, max_col=16):

            label = o_cell.value

            if label is None:
                continue

            label = str(label).strip()

            if label.upper() in _COMMODITY_HEADERS:
                current_commodity = _COMMODITY_HEADERS[label.upper()]
                continue

            if current_commodity is None:
                continue

            price = p_cell.value

            try:
                price = float(price)
            except (TypeError, ValueError):
                continue

            quotes.append({
                "commodity": current_commodity,
                "contract": _convert_cqg_symbol(label),
                "price": price,
            })

    return quotes
