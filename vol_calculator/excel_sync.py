from openpyxl.worksheet.datavalidation import DataValidation

from vol_calculator.instruments import Future, AmericanOption
from vol_calculator.greeks import GreekEngine
from vol_calculator.scenarios import spot_ladder, vol_ladder


def clear_range(ws, min_row=1, max_row=500, min_col=1, max_col=20):
    """Blanks out a rectangular range -- used before rewriting a sheet/block so a shrinking table doesn't leave stale cells behind."""

    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.value = None


def write_dataframe(ws, df, start_row=1, start_col=1, label=None):
    """
    Writes a pandas DataFrame as a plain header + rows block, starting at
    (start_row, start_col). If label is given, it goes on its own row
    just above the header. Returns the row just after the last one
    written, so callers can stack multiple tables on one sheet.
    """

    row = start_row

    if label is not None:
        ws.cell(row=row, column=start_col, value=label)
        row += 1

    for col, header in enumerate(df.columns, start=start_col):
        ws.cell(row=row, column=col, value=header)
    row += 1

    for _, data_row in df.iterrows():
        for col, value in enumerate(data_row, start=start_col):
            ws.cell(row=row, column=col, value=round(value, 4) if isinstance(value, float) else value)
        row += 1

    return row

POSITIONS_HEADERS = [
    "Contract", "Commodity", "Type", "Quantity",
    "Strike", "Expiry (d)", "Vol (%)", "Settlement Price",
]

FUTURE_SPOT_LABEL_CELL = "A10"
FUTURE_SPOT_VALUE_CELL = "B10"


def write_future_spot_selector(ws, contracts, label_cell=FUTURE_SPOT_LABEL_CELL, value_cell=FUTURE_SPOT_VALUE_CELL):
    """
    A native Excel dropdown (data validation list) of the contracts
    currently in the book -- the "which contract is shocked" selector
    for sheets that show one contract at a time (Ladders, Surface,
    Stress, Strike, Payoff, Breakevens), same role as the Future Spot
    selector in each Streamlit tab.

    The dropdown itself works live in Excel. What's NOT live yet (no
    xlwings): picking a new value here doesn't recompute the other
    sheets by itself -- the refresh script needs to be re-run after
    changing it, and it reads whatever's currently in value_cell.
    """

    ws[label_cell] = "Future Spot:"

    if not contracts:
        return

    current = ws[value_cell].value
    ws[value_cell] = current if current in contracts else contracts[0]

    dv = DataValidation(type="list", formula1=f'"{",".join(contracts)}"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[value_cell])


def read_future_spot(ws, contracts, value_cell=FUTURE_SPOT_VALUE_CELL):
    """Reads the current Future Spot selection, falling back to the first contract if unset/invalid."""

    value = ws[value_cell].value
    return value if value in contracts else (contracts[0] if contracts else None)


def contracts_from_imported(imported):
    """Distinct contracts, in order of first appearance."""

    seen = []
    for _, meta in imported:
        c = meta.get("contract")
        if c and c not in seen:
            seen.append(c)
    return seen


def prices_from_imported(imported):
    """
    contract -> settlement price, from the Daily Prop file itself --
    stand-in for live CQGXL prices until the Home sheet's formula cells
    have cached values to read (requires Excel to have actually
    recalculated with CQGXL connected and been saved at least once).
    """

    prices = {}
    for _, meta in imported:
        price = meta.get("settlement_price")
        if price is not None:
            prices[meta["contract"]] = price
    return prices


def _position_row(instrument, meta):

    if isinstance(instrument, Future):
        return [
            meta["contract"], meta["commodity"], "Future", instrument.quantity,
            None, None, meta.get("vol_pct"), meta.get("settlement_price"),
        ]

    if isinstance(instrument, AmericanOption):
        return [
            meta["contract"], meta["commodity"],
            "Call" if instrument.option_type == "call" else "Put",
            instrument.quantity, instrument.strike, instrument.expiry_days,
            round(instrument.volatility * 100, 2), meta.get("settlement_price"),
        ]

    raise TypeError(f"Unsupported instrument type: {type(instrument)}")


def write_positions_table(ws, imported, start_row=11, clear_rows=500):
    """
    Writes (instrument, meta) pairs -- the same shape import_bushel_positions
    returns -- into an existing worksheet, starting at start_row (a
    "POSITIONS" label, then headers, then one row per leg). Clears a
    generous range below start_row first, so a shrinking book doesn't
    leave stale rows behind from a previous, longer day's data.
    """

    n_cols = len(POSITIONS_HEADERS)

    for row in ws.iter_rows(
        min_row=start_row, max_row=start_row + clear_rows, min_col=1, max_col=n_cols
    ):
        for cell in row:
            cell.value = None

    ws.cell(row=start_row, column=1, value="POSITIONS")

    header_row = start_row + 1
    for col, header in enumerate(POSITIONS_HEADERS, start=1):
        ws.cell(row=header_row, column=col, value=header)

    for i, (instrument, meta) in enumerate(imported):
        row_values = _position_row(instrument, meta)
        for col, value in enumerate(row_values, start=1):
            ws.cell(row=header_row + 1 + i, column=col, value=value)

    return ws


def _round_greek(key, value):
    return round(value, 3) if key == "Value ($k)" else round(value, 2)


def write_greeks_sheet(ws, portfolio, market, contracts, clear_rows=500):
    """
    One row per contract, each computed with that contract as the shocked
    "Future Spot" -- same GreekEngine.report_dollars() used by the
    Streamlit Greeks tab, same decimal precision (Value ($k) at 3dp,
    every other Greek at 2dp). Gives all contracts at once instead of
    one-at-a-time like the Streamlit selector, since this is a static
    snapshot, not a live dropdown.
    """

    if not contracts:
        return ws

    sample = GreekEngine.report_dollars(portfolio, market, contracts[0])
    greek_names = list(sample.keys())
    n_cols = 1 + len(greek_names)

    for row in ws.iter_rows(min_row=1, max_row=clear_rows, min_col=1, max_col=n_cols):
        for cell in row:
            cell.value = None

    header = ["Contract"] + greek_names
    for col, value in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=value)

    for i, contract in enumerate(contracts):
        dollars = GreekEngine.report_dollars(portfolio, market, contract)
        row_values = [contract] + [_round_greek(k, v) for k, v in dollars.items()]
        for col, value in enumerate(row_values, start=1):
            ws.cell(row=2 + i, column=col, value=value)

    return ws


def write_spot_vol_ladder_sheet(ws, portfolio, market, contract, clear_rows=500, clear_cols=12):
    """
    Spot ladder table, then a gap, then vol ladder table below it --
    both for the one selected Future Spot contract, same shape/columns
    as the Streamlit Ladders tab.
    """

    clear_range(ws, min_row=1, max_row=clear_rows, min_col=1, max_col=clear_cols)

    spot_df = spot_ladder(portfolio, market, contract)
    next_row = write_dataframe(ws, spot_df, start_row=1, label="SPOT LADDER")

    vol_df = vol_ladder(portfolio, market, contract)
    write_dataframe(ws, vol_df, start_row=next_row + 1, label="VOL LADDER")

    return ws
